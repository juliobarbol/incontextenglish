import json
import subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# El banco vive en public/quiz.js, que es lo que se publica. Se lee de ahí y no
# de una copia en contenido/: dos copias se desincronizan sin que nadie lo note.
EXTRAER = """
const fs = require("fs");
const src = fs.readFileSync("public/quiz.js", "utf8");
const banco = eval(src.match(/const BANCO = (\\[[\\s\\S]*?\\n\\]);/)[1]);
process.stdout.write(JSON.stringify(banco));
"""
datos = json.loads(subprocess.run(["node", "-e", EXTRAER], capture_output=True,
                                  text=True, check=True).stdout)

LETRAS = ["A", "B", "C", "D"]

BORDO = "6E1112"
AMARILLO = "FBC93E"
CREMA = "F7F1E6"
GRIS = "F2F2F2"

fuente = "Arial"
titulo = Font(name=fuente, size=14, bold=True, color=BORDO)
cabecera = Font(name=fuente, size=10, bold=True, color="FFFFFF")
normal = Font(name=fuente, size=10)
mono = Font(name=fuente, size=10, color="444444")
chico = Font(name=fuente, size=9, color="666666")

relleno_cab = PatternFill("solid", fgColor=BORDO)
relleno_input = PatternFill("solid", fgColor=AMARILLO)
relleno_banda = PatternFill("solid", fgColor=CREMA)
relleno_gris = PatternFill("solid", fgColor=GRIS)

borde = Border(*[Side(style="thin", color="D9D9D9")] * 4)
arriba = Alignment(vertical="top", wrap_text=True)

wb = Workbook()

# ---------------------------------------------------------------- Cómo revisar
ins = wb.active
ins.title = "Cómo revisar"
ins.sheet_view.showGridLines = False

filas = [
    ("Test de nivel de In Context English — revisión del banco de preguntas", titulo),
    ("", None),
    ("Qué es esto", Font(name=fuente, size=11, bold=True)),
    ("El banco de preguntas del test de nivel del sitio, tal como está publicado hoy: 10 por cada "
     "banda (A1, A2, B1, B2 y C1). Cada persona que hace el test recibe 4 sorteadas de cada banda, "
     "así que dos personas no hacen el mismo test y no se puede memorizar repitiéndolo.", normal),
    ("", None),
    ("Esta planilla se genera leyendo el sitio, así que siempre muestra lo que está en vivo. "
     "Nada cambia hasta que devuelvas el archivo.",
     Font(name=fuente, size=10, bold=True, color=BORDO)),
    ("", None),
    ("Cómo están escritas", Font(name=fuente, size=11, bold=True)),
    ("Cada pregunta es un intercambio corto de dos líneas, casi siempre en una situación de "
     "trabajo: una reunión, un mail, una llamada. La idea es que el test se parezca al método que "
     "enseñás, en vez de tomar gramática aislada.", normal),
    ("", None),
    ("Qué te pedimos", Font(name=fuente, size=11, bold=True)),
    ("En la hoja «Preguntas», completá las dos últimas columnas, que están pintadas de amarillo. "
     "El resto es sólo para leer.", normal),
    ("", None),
    ("Columna «¿Va?» — elegí una opción de la lista desplegable:", Font(name=fuente, size=10, bold=True)),
    ("      Sí          la pregunta queda como está", normal),
    ("      Cambiar     sirve la idea pero hay que corregir algo (decilo en el comentario)", normal),
    ("      Sacar       no sirve para esa banda o directamente no va", normal),
    ("", None),
    ("Columna «Comentario» — escribí lo que quieras: si la opción marcada como correcta no lo es, "
     "si hay dos que podrían serlo, si la pregunta es más fácil o más difícil de lo que dice su "
     "banda, si el inglés suena raro, o una redacción mejor.", normal),
    ("", None),
    ("Así se ve una fila completada", Font(name=fuente, size=11, bold=True)),
]

f = 1
for texto, estilo in filas:
    c = ins.cell(row=f, column=1, value=texto)
    if estilo:
        c.font = estilo
    c.alignment = Alignment(vertical="top", wrap_text=True)
    ins.row_dimensions[f].height = 42 if (estilo is normal and len(texto) > 110) else 15
    f += 1

ejemplo_cab = ["Nº", "Banda", "Pregunta", "Correcta", "¿Va?", "Comentario"]
ejemplo_val = [
    17, "A2", "— Is the new system better?\n— It's ___ than the old one, yes.", "A (faster)",
    "Cambiar", "Para A2 pondría «easier to use» en lugar de «faster»: se usa más.",
]
f += 1
for col, (h, v) in enumerate(zip(ejemplo_cab, ejemplo_val), start=1):
    ch = ins.cell(row=f, column=col, value=h)
    ch.font = cabecera
    ch.fill = relleno_cab
    ch.alignment = Alignment(vertical="center", horizontal="center")
    cv = ins.cell(row=f + 1, column=col, value=v)
    cv.font = normal
    cv.alignment = arriba
    cv.border = borde
    if h in ("¿Va?", "Comentario"):
        cv.fill = relleno_input
ins.row_dimensions[f + 1].height = 46

for col, ancho in zip("ABCDEF", [10, 9, 44, 16, 13, 52]):
    ins.column_dimensions[col].width = ancho
ins.column_dimensions["A"].width = 62

f += 4
nota = ins.cell(row=f, column=1,
                value="Cuando termines, guardá el archivo y devolvelo. Con eso se arma la versión "
                      "final y recién ahí se publica en el sitio.")
nota.font = chico
nota.alignment = Alignment(vertical="top", wrap_text=True)

# ------------------------------------------------------------------- Preguntas
ws = wb.create_sheet("Preguntas")
ws.sheet_view.showGridLines = False

cols = [
    ("Nº", 6), ("Banda", 8), ("Foco gramatical", 26), ("Pregunta", 54),
    ("Opción A", 17), ("Opción B", 17), ("Opción C", 17), ("Opción D", 17),
    ("Correcta", 18), ("¿Va?", 12), ("Comentario", 48),
]
for i, (nombre, ancho) in enumerate(cols, start=1):
    c = ws.cell(row=1, column=i, value=nombre)
    c.font = cabecera
    c.fill = relleno_cab
    c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = ancho
ws.row_dimensions[1].height = 28

for j, p in enumerate(datos):
    fila = j + 2
    valores = [
        j + 1, p["banda"], p["foco"], p["q"],
        p["opts"][0], p["opts"][1], p["opts"][2], p["opts"][3],
        f'{LETRAS[p["a"]]} ({p["opts"][p["a"]]})', None, None,
    ]
    for i, v in enumerate(valores, start=1):
        c = ws.cell(row=fila, column=i, value=v)
        c.font = normal
        c.alignment = arriba
        c.border = borde
        if i == 2:
            c.fill = relleno_banda
            c.alignment = Alignment(vertical="top", horizontal="center")
            c.font = Font(name=fuente, size=10, bold=True, color=BORDO)
        if i == 4:
            c.font = mono
        if i == 9:
            c.font = Font(name=fuente, size=10, bold=True)
        if i in (10, 11):
            c.fill = relleno_input
    # Sin alto fijo: las preguntas son de dos líneas y algunas envuelven a tres,
    # así que conviene que Excel ajuste solo antes que recortar el texto.

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:K{len(datos) + 1}"

dv = DataValidation(type="list", formula1='"Sí,Cambiar,Sacar"', allow_blank=True, showDropDown=False)
dv.prompt = "Sí / Cambiar / Sacar"
dv.promptTitle = "¿Va esta pregunta?"
ws.add_data_validation(dv)
dv.add(f"J2:J{len(datos) + 1}")

# --------------------------------------------------------------------- Resumen
res = wb.create_sheet("Resumen")
res.sheet_view.showGridLines = False
ultima = len(datos) + 1

res["A1"] = "Resumen de la revisión"
res["A1"].font = titulo
res["A3"] = "Se actualiza solo a medida que completás la columna «¿Va?»."
res["A3"].font = chico

enc = ["Banda", "Preguntas", "Sí", "Cambiar", "Sacar", "Sin revisar"]
for i, h in enumerate(enc, start=1):
    c = res.cell(row=5, column=i, value=h)
    c.font = cabecera
    c.fill = relleno_cab
    c.alignment = Alignment(horizontal="center", vertical="center")
    res.column_dimensions[get_column_letter(i)].width = 14

for k, banda in enumerate(["A1", "A2", "B1", "B2", "C1"]):
    r = 6 + k
    res.cell(row=r, column=1, value=banda).font = Font(name=fuente, size=10, bold=True, color=BORDO)
    res.cell(row=r, column=2, value=f'=COUNTIF(Preguntas!$B$2:$B${ultima},$A{r})')
    for i, veredicto in enumerate(["Sí", "Cambiar", "Sacar"], start=3):
        res.cell(row=r, column=i,
                 value=f'=COUNTIFS(Preguntas!$B$2:$B${ultima},$A{r},Preguntas!$J$2:$J${ultima},"{veredicto}")')
    res.cell(row=r, column=6, value=f"=$B{r}-SUM($C{r}:$E{r})")
    for i in range(1, 7):
        cc = res.cell(row=r, column=i)
        cc.border = borde
        if i > 1:
            cc.font = normal
            cc.alignment = Alignment(horizontal="center")
        else:
            cc.alignment = Alignment(horizontal="center")

r = 11
res.cell(row=r, column=1, value="Total").font = Font(name=fuente, size=10, bold=True)
for i in range(2, 7):
    c = res.cell(row=r, column=i, value=f"=SUM({get_column_letter(i)}6:{get_column_letter(i)}10)")
    c.font = Font(name=fuente, size=10, bold=True)
    c.alignment = Alignment(horizontal="center")
    c.fill = relleno_gris
for i in range(1, 7):
    res.cell(row=r, column=i).border = borde

res["A13"] = "Objetivo: al menos 8 preguntas utilizables por banda, para que el sorteo de 4 tenga de dónde elegir."
res["A13"].font = chico

SALIDA = "contenido/test-de-nivel-preguntas-para-revisar.xlsx"
wb.calculation.fullCalcOnLoad = True
wb.save(SALIDA)

# ---------------------------------------------------------------------------
# openpyxl escribe las fórmulas sin valor en caché, así que hasta que alguien
# abra el archivo en Excel las celdas se ven vacías en cualquier vista previa.
# Como el archivo se entrega con la columna «¿Va?» en blanco, los valores de
# partida se calculan acá y se inyectan como caché. Al abrirlo, Excel recalcula
# igual (fullCalcOnLoad) y siguen actualizándose a medida que se completa.
# ---------------------------------------------------------------------------
import re, shutil, zipfile

por_banda = {b: sum(1 for p in datos if p["banda"] == b) for b in ["A1", "A2", "B1", "B2", "C1"]}
cache = {}
for k, banda in enumerate(["A1", "A2", "B1", "B2", "C1"]):
    r = 6 + k
    cache[f"B{r}"] = por_banda[banda]      # preguntas de la banda
    cache[f"C{r}"] = cache[f"D{r}"] = cache[f"E{r}"] = 0   # sí / cambiar / sacar
    cache[f"F{r}"] = por_banda[banda]      # sin revisar
cache["B11"] = sum(por_banda.values())
cache["C11"] = cache["D11"] = cache["E11"] = 0
cache["F11"] = sum(por_banda.values())

hoja_resumen = "xl/worksheets/sheet%d.xml" % (wb.sheetnames.index("Resumen") + 1)
tmp = SALIDA + ".tmp"
with zipfile.ZipFile(SALIDA) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        contenido = zin.read(item.filename)
        if item.filename == hoja_resumen:
            xml = contenido.decode("utf8")
            for ref, valor in cache.items():
                xml, n = re.subn(
                    r'(<c r="%s"[^>]*>)(<f>.*?</f>)(?:<v\s*/>|<v>.*?</v>)?(</c>)' % ref,
                    r'\1\2<v>%s</v>\3' % valor,
                    xml,
                )
                if n != 1:
                    raise SystemExit(f"no pude escribir el valor de {ref} (coincidencias: {n})")
            contenido = xml.encode("utf8")
        zout.writestr(item, contenido)
shutil.move(tmp, SALIDA)
print("planilla escrita, con", len(cache), "valores en caché")
